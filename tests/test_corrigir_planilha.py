"""Write-back resiliente: uma célula com problema é pulada e as demais correções são gravadas.

Planilhas criadas em pasta temporária com valores fictícios.
"""

import hashlib

import openpyxl
import pytest
from lxml import etree

from corrigir_planilha import NS, _patch_aba, corrigir_planilha, nome_da_copia
from localizar_arquivos import localizar_arquivo

ABA = "UTI PEDIATRICA SET"


def _planilha(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append(["N. DE SAIDAS", 1, 2, 3, "=B1+C1"])  # TOTAL quebrado: esqueceu o dia 3
    ws.append(["TAXA FICTICIA", None, None, None, "=E1/100"])
    caminho = tmp_path / "PRODUCAO FICTICIA.xlsx"
    wb.save(caminho)
    return caminho


def _meta(rotulo_taxa="TAXA FICTICIA"):
    celula = {"col_valor": 5, "col_ini": 2, "col_fim": 4}
    return {"aba": ABA, "col_rotulo": 1, "celulas": {
        "saidas": {"linha": 1, "rotulo": "N. DE SAIDAS", **celula},
        "taxa": {"linha": 2, "rotulo": rotulo_taxa, **celula},
    }}


def _hash(caminho):
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def test_celula_com_problema_e_pulada_e_as_demais_sao_gravadas(tmp_path):
    origem = _planilha(tmp_path)
    antes = _hash(origem)
    correcoes = [
        {"chave": "saidas", "novo": 6.0, "acao": ("soma_linha",)},
        {"chave": "taxa", "novo": 0.5, "acao": ("div", "saidas", "indicador_nao_rastreado")},
    ]
    destino, edicoes, puladas = corrigir_planilha(origem, _meta(), correcoes, "SETEMBRO")

    assert destino == tmp_path / "PRODUCAO FICTICIA_Corrigida_SETEMBRO.xlsx" and destino.is_file()
    assert [e["celula"] for e in edicoes] == [f"{ABA}!E1"]
    assert [(p["chave"], p["celula"]) for p in puladas] == [("taxa", f"{ABA}!E2")]
    assert "indicador_nao_rastreado" in puladas[0]["motivo"]

    ws = openpyxl.load_workbook(destino)[ABA]
    assert ws["E1"].value == "=SUM(B1:D1)"
    assert ws["E2"].value == "=E1/100"  # pulada: intocada
    assert _hash(origem) == antes  # original nunca é alterado
    assert not list(tmp_path.glob("*.tmp"))  # temporário removido


def test_rotulo_diferente_pula_so_aquela_celula(tmp_path):
    origem = _planilha(tmp_path)
    correcoes = [
        {"chave": "saidas", "novo": 6.0, "acao": ("soma_linha",)},
        {"chave": "taxa", "novo": 0.06, "acao": ("div", "saidas", "saidas")},
    ]
    _, edicoes, puladas = corrigir_planilha(origem, _meta(rotulo_taxa="OUTRO ROTULO"), correcoes, "SETEMBRO")
    assert len(edicoes) == 1
    assert "rótulo" in puladas[0]["motivo"]


def test_sem_nenhuma_correcao_aplicavel_nao_grava_arquivo(tmp_path):
    origem = _planilha(tmp_path)
    correcoes = [{"chave": "taxa", "novo": 0.5, "acao": ("div", "saidas", "indicador_nao_rastreado")}]
    destino, edicoes, puladas = corrigir_planilha(origem, _meta(), correcoes, "SETEMBRO")
    assert destino is None and edicoes == [] and len(puladas) == 1
    assert not list(tmp_path.glob("*_Corrigida*"))


@pytest.mark.parametrize("mes", ["set", "9", "Setembro"])
def test_nome_da_copia_leva_o_mes_normalizado(mes):
    assert nome_da_copia("PLANILHA INDICADORES.xlsx", mes) == "PLANILHA INDICADORES_Corrigida_SETEMBRO.xlsx"


def test_rodar_outro_mes_nao_sobrescreve_a_copia_anterior(tmp_path):
    origem = _planilha(tmp_path)
    correcoes = [{"chave": "saidas", "novo": 6.0, "acao": ("soma_linha",)}]
    setembro, _, _ = corrigir_planilha(origem, _meta(), correcoes, "SETEMBRO")
    antes = _hash(setembro)
    julho, _, _ = corrigir_planilha(origem, _meta(), correcoes, "JULHO")

    assert julho.name == "PRODUCAO FICTICIA_Corrigida_JULHO.xlsx"
    assert setembro.is_file() and _hash(setembro) == antes


def test_copias_corrigidas_nunca_sao_lidas_como_fonte(tmp_path):
    """A cópia cita o mês no nome e ganharia do original anual na escolha por mês se não fosse excluída."""
    (tmp_path / "PRODUCAO DIARIA 2026.xlsx").write_text("x")
    (tmp_path / "PRODUCAO DIARIA 2026_Corrigida_SETEMBRO.xlsx").write_text("x")
    (tmp_path / "PRODUCAO DIARIA 2026_Corrigida.xlsx").write_text("x")  # formato antigo
    escolhido = localizar_arquivo(tmp_path, ".xlsx", ["produ"], "SETEMBRO", 2026)
    assert escolhido.name == "PRODUCAO DIARIA 2026.xlsx"


XML_ABA = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="{NS}"><sheetData><row r="1">
<c r="A1"><f t="shared" ref="A1:A3" si="0">B1*2</f><v>2</v></c>
<c r="B1"><v>1</v></c>
<c r="C1"><f>B1/100</f><v>0.01</v></c>
</row></sheetData></worksheet>""".encode()


def _edicao(formula, valor):
    return {"chave": "x", "atual": None, "formula": formula, "valor": valor, "mesma": False}


def test_patch_xml_pula_formula_compartilhada_e_celula_inexistente():
    edicoes = {"A1": _edicao("=B1*3", 3.0), "C1": _edicao("=B1/B1", 1.0), "Z9": _edicao("=1", 1.0)}
    xml, puladas = _patch_aba(XML_ABA, edicoes)

    assert set(puladas) == {"A1", "Z9"}
    assert "compartilhada" in puladas["A1"]
    raiz = etree.fromstring(xml)
    celula = {c.get("r"): c for c in raiz.iter(f"{{{NS}}}c")}
    a1_f = celula["A1"].find(f"{{{NS}}}f")
    assert (a1_f.text, a1_f.get("t"), celula["A1"].findtext(f"{{{NS}}}v")) == ("B1*2", "shared", "2")
    assert celula["C1"].findtext(f"{{{NS}}}f") == "B1/B1"
    assert celula["C1"].findtext(f"{{{NS}}}v") == "1.0"
