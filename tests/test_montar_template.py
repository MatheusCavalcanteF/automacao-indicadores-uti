"""Parágrafos cujo texto vem do documento-base: só o número vira variável. Documentos sintéticos."""

from docx import Document

from gerador_relatorio import gerar_relatorio_word
from montar_template import montar


def _base(tmp_path, *linhas):
    doc = Document()
    for linha in linhas:
        doc.add_paragraph(linha)
    caminho = tmp_path / "base.docx"
    doc.save(caminho)
    return caminho


def _textos(caminho):
    return [p.text for p in Document(caminho).paragraphs]


def test_nascidos_vivos_mantem_o_texto_do_documento_base(tmp_path):
    base = _base(tmp_path, "Tivemos um total de 57 nascidos vivos na Unidade Exemplo.")
    modelo = montar(base, tmp_path / "modelo.docx")

    assert _textos(modelo) == ["Tivemos um total de {{ n_de_nascidos_vivos | num }} nascidos vivos na Unidade Exemplo."]

    saida, faltantes, _ = gerar_relatorio_word({"n_de_nascidos_vivos": 61.0}, modelo, tmp_path / "r.docx")
    assert faltantes == []
    assert _textos(saida)[-1] == "Tivemos um total de 61 nascidos vivos na Unidade Exemplo."


def test_numero_com_separador_de_milhar_vira_uma_variavel_so(tmp_path):
    base = _base(tmp_path, "Tivemos um total de 1.234 nascidos vivos (Unidade Exemplo), 2 a mais que o mês anterior.")
    modelo = montar(base, tmp_path / "modelo.docx")

    assert _textos(modelo) == [
        "Tivemos um total de {{ n_de_nascidos_vivos | num }} nascidos vivos (Unidade Exemplo), 2 a mais que o mês anterior."
    ]


def test_paragrafo_sem_numero_fica_intacto_e_e_avisado(tmp_path, capsys):
    texto = "Tivemos um total de vários nascidos vivos na Unidade Exemplo."
    modelo = montar(_base(tmp_path, texto), tmp_path / "modelo.docx")

    assert _textos(modelo) == [texto]
    avisos = capsys.readouterr().out
    assert "parágrafos de base não encontrados" in avisos and "Tivemos um total de" in avisos
