"""Ponta a ponta com os dados fictícios de scripts/gerar_dados_exemplo.py, tudo em tmp_path.

O modelo Word também é o de exemplo (criar_template_exemplo): o modelo real em templates/ nunca é usado.
"""

import hashlib
import re

import openpyxl
import pytest
from docx import Document

import main
from gerar_dados_exemplo import PASTA_REAL, gerar


@pytest.fixture
def pastas(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "TEMPLATE", tmp_path / "templates" / "modelo_relatorio.docx")
    entrada, saida = tmp_path / "entrada", tmp_path / "saida"
    producao, censo = gerar(entrada, "setembro", 2026)
    return entrada, saida, producao, censo


def _rodar(entrada, saida, *extra):
    main.main(["setembro", "2026", "--pasta-entrada", str(entrada), "--pasta-saida", str(saida), *extra])


def _sha(caminho):
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def test_gera_relatorio_copia_corrigida_e_log(pastas):
    entrada, saida, producao, censo = pastas
    antes = _sha(producao)

    _rodar(entrada, saida)

    textos = [p.text for p in Document(saida / "Relatorio_Setembro_2026_Final.docx").paragraphs]
    assert textos[0] == "AVISOS CRÍTICOS E PENDÊNCIAS"  # PDF e ODT ausentes
    assert any("FISIOTERAPIA" in t for t in textos)
    # Mortalidade veio x100 na planilha e foi recalculada: sai como percentual de um dígito, não 303%.
    assert any(re.fullmatch(r"taxa_de_mortalidade_em_uti: \d,\d\d%", t) for t in textos)
    assert "censo_paciente_dia: [Sem Dados]" not in textos

    copia = entrada / "PLANILHA PRODUCAO DIARIA EXEMPLO 2026_Corrigida_SETEMBRO.xlsx"
    formulas = [c.value for linha in openpyxl.load_workbook(copia).active.iter_rows() for c in linha
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(") for f in formulas) and len(formulas) == 3
    assert _sha(producao) == antes  # original nunca é alterado

    (log,) = saida.glob("execucao_SETEMBRO_2026_*.log")
    conteudo = log.read_text(encoding="utf-8")
    assert _sha(producao) in conteudo and _sha(censo) in conteudo
    assert "código de saída 0" in conteudo
    assert (saida / "mapa_variaveis.txt").is_file()
    assert main.TEMPLATE.is_file()  # modelo de exemplo criado no lugar do monkeypatch


def test_estrito_bloqueia_por_fonte_ausente(pastas):
    entrada, saida, _, _ = pastas
    with pytest.raises(SystemExit) as saida_estrito:
        _rodar(entrada, saida, "--estrito")
    assert saida_estrito.value.code == main.SAIDA_BLOQUEADA
    assert not list(saida.glob("*.docx"))


def test_sem_planilha_de_producao_falha(tmp_path):
    with pytest.raises(SystemExit) as falha:
        _rodar(tmp_path / "vazia", tmp_path / "saida")
    assert falha.value.code == main.SAIDA_FALHA


def test_gerador_e_deterministico_e_nao_escreve_na_pasta_real(tmp_path):
    a, b = gerar(tmp_path / "a"), gerar(tmp_path / "b")
    linhas = [[c.value for c in r] for r in openpyxl.load_workbook(a[0]).active.iter_rows()]
    assert linhas == [[c.value for c in r] for r in openpyxl.load_workbook(b[0]).active.iter_rows()]
    with pytest.raises(ValueError, match="dados reais"):
        gerar(PASTA_REAL)
