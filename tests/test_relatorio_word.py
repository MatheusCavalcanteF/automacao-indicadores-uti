"""Dia 31 nas tabelas diárias e página de avisos no Word, com documentos sintéticos."""

from docx import Document

from gerador_relatorio import dia, dias_no_modelo, gerar_relatorio_word, inserir_pagina_avisos, montar_contexto
from montar_template import DIAS_MAX, garantir_colunas_de_dia


def _tabela_diaria(doc, dias=30):
    """rótulo | dias | TOTAL | taxa, com a 1ª linha mesclada sobre rótulo + dias (como no modelo-base)."""
    tabela = doc.add_table(rows=3, cols=dias + 3)
    tabela.cell(0, 0).merge(tabela.cell(0, dias)).text = "TITULO"
    tabela.cell(1, 0).text = "N. FICTICIO"
    for i in range(1, dias + 1):
        tabela.cell(1, i).text = str(i)
    return tabela


def _grade_consistente(tabela):
    total = len(tabela._tbl.tblGrid.gridCol_lst)
    return all(sum(tc.grid_span for tc in tr.tc_lst) == total for tr in tabela._tbl.tr_lst)


def test_tabela_de_30_dias_ganha_a_coluna_31_vazia():
    doc = Document()
    tabela = _tabela_diaria(doc)
    largura_antes = sum(int(c.w) for c in tabela._tbl.tblGrid.gridCol_lst)

    assert garantir_colunas_de_dia(tabela) == 1
    assert len(tabela._tbl.tblGrid.gridCol_lst) == DIAS_MAX + 3
    assert _grade_consistente(tabela)
    assert tabela._tbl.tr_lst[0].tc_lst[0].grid_span == DIAS_MAX + 1  # título cobre o dia 31
    linha = [c.text for c in tabela.rows[1].cells]
    assert linha[30:32] == ["30", ""]  # cópia da coluna 30, sem o texto
    assert sum(int(c.w) for c in tabela._tbl.tblGrid.gridCol_lst) == largura_antes


def test_ampliacao_e_idempotente():
    tabela = _tabela_diaria(Document())
    garantir_colunas_de_dia(tabela)
    assert garantir_colunas_de_dia(tabela) == 0
    assert len(tabela._tbl.tblGrid.gridCol_lst) == DIAS_MAX + 3


def test_tabela_de_28_dias_ganha_tres_colunas():
    tabela = _tabela_diaria(Document(), dias=28)
    assert garantir_colunas_de_dia(tabela) == 3
    assert _grade_consistente(tabela)


def test_serie_e_cortada_nos_dias_do_mes():
    serie = [1.0] * 31  # aba com uma coluna a mais que o mês
    contexto = montar_contexto({}, {"n_dias": 30}, {"x": serie})
    assert len(contexto["serie__x"]) == 30
    assert dia(contexto["serie__x"], 31) == ""
    assert dia(contexto["serie__x"], 30) == "1"


def test_dia_vazio_para_dias_inexistentes_ou_sem_valor():
    serie = [5.0, None, 2.5]
    assert [dia(serie, i) for i in (0, 1, 2, 3, 4)] == ["", "5", "", "2,50", ""]


def test_pagina_de_avisos_vai_para_o_inicio_e_quebra_pagina():
    doc = Document()
    doc.add_paragraph("CAPA DO RELATORIO")
    inserir_pagina_avisos(doc, {"CRITICO": ["taxa X inválida"], "FONTE": ["PDF ausente"], "CONFERIR": []}, "SETEMBRO/2026")

    textos = [p.text for p in doc.paragraphs]
    assert textos[0] == "AVISOS CRÍTICOS E PENDÊNCIAS"
    assert "SETEMBRO/2026" in textos[1]
    assert any("taxa X inválida" in t for t in textos) and any("PDF ausente" in t for t in textos)
    assert not any("Divergências" in t for t in textos)  # seção vazia não aparece
    assert textos[-1] == "CAPA DO RELATORIO"
    assert 'w:type="page"' in doc.paragraphs[-2]._p.xml
    assert not any(p.style.name.startswith("Heading") for p in doc.paragraphs[:-1])  # fora do sumário


def _modelo(tmp_path, *linhas):
    doc = Document()
    for linha in linhas:
        doc.add_paragraph(linha)
    caminho = tmp_path / "modelo.docx"
    doc.save(caminho)
    return caminho


def test_relatorio_sem_pendencias_nao_tem_pagina_de_avisos(tmp_path):
    modelo = _modelo(tmp_path, "Saídas: {{ saidas | num }}")
    saida, faltantes, com_pagina = gerar_relatorio_word({"saidas": 12.0}, modelo, tmp_path / "r.docx",
                                                        avisos={"CONFERIR": ["só uma divergência"]})
    assert (faltantes, com_pagina) == ([], False)
    assert [p.text for p in Document(saida).paragraphs] == ["Saídas: 12"]


def test_variavel_faltante_gera_pagina_de_avisos(tmp_path):
    modelo = _modelo(tmp_path, "Saídas: {{ saidas | num }}", "Óbitos: {{ obitos | num }}")
    saida, faltantes, com_pagina = gerar_relatorio_word({"saidas": 12.0}, modelo, tmp_path / "r.docx")
    textos = [p.text for p in Document(saida).paragraphs]
    assert faltantes == ["obitos"] and com_pagina
    assert textos[0] == "AVISOS CRÍTICOS E PENDÊNCIAS"
    assert any("obitos" in t for t in textos[:-2])
    assert textos[-1] == "Óbitos: [Sem Dados]"


def test_dias_no_modelo(tmp_path):
    assert dias_no_modelo(_modelo(tmp_path, "{{ serie__x | dia(1) }} {{ serie__x | dia(31) }}")) == 31
    assert dias_no_modelo(_modelo(tmp_path, "sem tabela diária")) is None
    assert dias_no_modelo(tmp_path / "nao_existe.docx") is None
