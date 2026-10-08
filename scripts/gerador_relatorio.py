"""Preenche o modelo Word do Relatório de Avaliação dos Indicadores com docxtpl.

O contexto leva números crus; a formatação é feita no template por filtros Jinja:
  {{ taxa_de_ocupacao_global | pct }}   -> 93,67%
  {{ n_de_saidas | num }}               -> 29 (inteiro sem casas; senão 2 casas, vírgula)
  {{ media_de_permanencia | dec }}      -> 7,80
  {{ serie__x | dia(3) }}               -> valor do 3º dia da série 'x'
Variável inexistente não quebra a geração: aparece como [Sem Dados] e é listada.
Valores já pendentes (strings) passam pelos filtros sem alteração.

Se houver avisos críticos ou pendências, o documento ganha uma primeira página
"AVISOS CRÍTICOS E PENDÊNCIAS" listando-os (ver inserir_pagina_avisos).
"""

import re
import zipfile
from pathlib import Path

import jinja2
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from docxtpl import DocxTemplate

from extracao_producao import chave_variavel
from utils import eh_numero as _numerico, periodo_da_cli

RAIZ = Path(__file__).resolve().parent.parent

chave_template = chave_variavel


SEM_DADOS = "[Sem Dados]"

# Seções da página de avisos: (categoria, título). Só CRITICO, FONTE e PENDENCIA fazem a página aparecer;
# CONFERIR entra junto quando a página existe.
SECOES_AVISOS = (
    ("CRITICO", "Erros não corrigidos automaticamente"),
    ("FONTE", "Fontes de dados ausentes"),
    ("PENDENCIA", "Pendências de preenchimento"),
    ("CONFERIR", "Divergências a conferir"),
)
CATEGORIAS_QUE_EXIGEM_PAGINA = ("CRITICO", "FONTE", "PENDENCIA")
VERMELHO = RGBColor(0xC0, 0x00, 0x00)


class _Faltante(jinja2.Undefined):
    """Undefined que registra o nome da variável e deixa um marcador visível no Word."""

    registro = set()

    def __str__(self):
        _Faltante.registro.add(self._undefined_name)
        return SEM_DADOS

    __html__ = __str__


def _br(valor, casas):
    return f"{valor:.{casas}f}".replace(".", ",")


def _texto(valor):
    """Não numérico (marcador de pendência, variável ausente, None) passa direto, sem quebrar."""
    return "-" if valor is None else str(valor)


def pct(valor, casas=2):
    """0.9367 -> '93,67%'."""
    return _br(valor * 100, casas) + "%" if _numerico(valor) else _texto(valor)


def dec(valor, casas=2):
    return _br(valor, casas) if _numerico(valor) else _texto(valor)


def num(valor):
    """Inteiro sem casas; fracionário com 2 casas."""
    if not _numerico(valor):
        return _texto(valor)
    return str(int(valor)) if float(valor).is_integer() else _br(valor, 2)


def dia(serie, indice):
    """Valor do dia `indice` (1 = primeiro dia) de uma série; vazio se o mês não tem esse dia."""
    if not isinstance(serie, (list, tuple)) or not 1 <= indice <= len(serie) or serie[indice - 1] is None:
        return ""
    return num(serie[indice - 1])


def _ambiente():
    env = jinja2.Environment(undefined=_Faltante, autoescape=False)
    env.filters.update(pct=pct, dec=dec, num=num, fmt=num, dia=dia)
    return env


def montar_contexto(dados, extras=None, series=None):
    """{variável: número} + extras (mês, ano...) + séries diárias como 'serie__<variável>'.

    Com extras["n_dias"], cada série é cortada nesse número de dias: as colunas de dia que o mês
    não tem (29–31) saem em branco no Word, mesmo que a aba traga colunas a mais.
    """
    extras = extras or {}
    n_dias = extras.get("n_dias")
    contexto = dict(dados)
    for chave, valores in (series or {}).items():
        contexto[f"serie__{chave}"] = list(valores[:n_dias]) if n_dias else valores
    contexto.update(extras)
    return contexto


def dias_no_modelo(caminho_template):
    """Maior N usado em `dia(N)` no modelo (= colunas de dia das tabelas diárias), ou None se não houver."""
    try:
        with zipfile.ZipFile(caminho_template) as z:
            xml = z.read("word/document.xml").decode("utf-8")
    except (OSError, KeyError, zipfile.BadZipFile):
        return None
    dias = [int(n) for n in re.findall(r"dia\(\s*(\d+)\s*\)", xml)]
    return max(dias) if dias else None


def criar_template_exemplo(caminho_template, dados):
    """Modelo mínimo com todas as variáveis, só para testes quando não há modelo real."""
    doc = Document()
    doc.add_heading("Relatório de Avaliação dos Indicadores - UTI Pediátrica (MODELO DE EXEMPLO)", level=1)
    doc.add_paragraph("Referência: {{ mes }}/{{ ano }}")
    for nome in dados:
        filtro = "pct" if nome.startswith("taxa") and "__" not in nome else "num"
        doc.add_paragraph(f"{nome}: {{{{ {nome} | {filtro} }}}}")
    Path(caminho_template).parent.mkdir(parents=True, exist_ok=True)
    doc.save(caminho_template)


def exige_pagina_avisos(avisos):
    return any(avisos.get(c) for c in CATEGORIAS_QUE_EXIGEM_PAGINA)


def inserir_pagina_avisos(doc, avisos, referencia=""):
    """Insere no início de `doc` (python-docx) a página "AVISOS CRÍTICOS E PENDÊNCIAS".

    avisos: {categoria: [mensagens]} com as categorias de SECOES_AVISOS. Usa formatação direta, não
    estilos de título, para a página não entrar no sumário do relatório.
    """
    paragrafos = []

    def linha(texto, tamanho=11, negrito=False, cor=None):
        p = doc.add_paragraph()
        run = p.add_run(texto)
        run.bold, run.font.size = negrito, Pt(tamanho)
        if cor is not None:
            run.font.color.rgb = cor
        paragrafos.append(p)
        return p

    linha("AVISOS CRÍTICOS E PENDÊNCIAS", 16, True, VERMELHO)
    linha(f"Relatório gerado automaticamente{' — ' + referencia if referencia else ''}. Resolva os itens "
          "abaixo e remova esta página antes de enviar o relatório.", 10)
    for categoria, titulo in SECOES_AVISOS:
        itens = avisos.get(categoria) or []
        if not itens:
            continue
        linha(f"{titulo} ({len(itens)})", 12, True, VERMELHO if categoria in ("CRITICO", "FONTE") else None)
        for item in itens:
            linha(f"•  {item}", 10)
    paragrafos[-1].add_run().add_break(WD_BREAK.PAGE)

    corpo = doc.element.body
    for i, p in enumerate(paragrafos):
        corpo.insert(i, p._p)  # add_paragraph anexou no fim; move para o topo, na ordem


def gerar_relatorio_word(dados, caminho_template, caminho_saida, extras=None, series=None, avisos=None):
    """Renderiza o template e salva em caminho_saida.

    avisos: {categoria: [mensagens]} (ver SECOES_AVISOS). As variáveis do modelo sem valor entram como
    PENDENCIA; se houver CRITICO, FONTE ou PENDENCIA, a página de avisos é inserida no início.
    Retorna (Path salvo, variáveis faltantes, página de avisos inserida?).
    """
    caminho_template = Path(caminho_template)
    if not caminho_template.is_file():
        raise FileNotFoundError(f"Template não encontrado: {caminho_template}")

    _Faltante.registro.clear()
    try:
        modelo = DocxTemplate(caminho_template)
        modelo.render(montar_contexto(dados, extras, series), _ambiente())
    except Exception as e:
        raise ValueError(f"Falha ao renderizar '{caminho_template.name}': {e}") from e
    faltantes = sorted(_Faltante.registro)

    avisos = {c: list(v) for c, v in (avisos or {}).items()}
    if faltantes:
        avisos.setdefault("PENDENCIA", []).append(
            f"{len(faltantes)} variável(is) do modelo sem valor (aparecem como '{SEM_DADOS}'): {', '.join(faltantes)}."
        )
    com_pagina = exige_pagina_avisos(avisos)
    if com_pagina:
        extras = extras or {}
        referencia = f"{extras['mes']}/{extras['ano']}" if "mes" in extras and "ano" in extras else ""
        inserir_pagina_avisos(modelo.docx, avisos, referencia)

    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    try:
        modelo.save(caminho_saida)
    except PermissionError as e:
        raise PermissionError(f"Feche o arquivo se estiver aberto no Word: {caminho_saida}") from e
    return caminho_saida, faltantes, com_pagina


if __name__ == "__main__":
    import sys

    import log_execucao
    from consolidacao import consolidar_e_sanitizar, extrair_por_fonte

    log_execucao.configurar()
    periodo = periodo_da_cli(sys.argv[1:], "python scripts/gerador_relatorio.py MES [ANO]")
    secoes, auxiliar = extrair_por_fonte(periodo.mes, periodo.ano)
    dados, _, _ = consolidar_e_sanitizar(secoes, auxiliar["series"])
    template = RAIZ / "templates" / "modelo_relatorio.docx"
    saida = RAIZ / "relatorios_gerados" / f"Relatorio_{periodo.mes.capitalize()}_{periodo.ano}_Teste.docx"
    arquivo, faltantes, _ = gerar_relatorio_word(
        dados, template, saida, extras={"mes": periodo.mes, "ano": periodo.ano, "n_dias": periodo.n_dias},
        series=auxiliar["series"],
    )
    print("Relatório salvo em:", arquivo, "| variáveis ausentes:", faltantes)
