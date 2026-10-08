"""Monta templates/modelo_relatorio.docx a partir do relatório de um mês anterior (.docx).

Troca os números do mês por variáveis docxtpl ({{ var | filtro }}), esvazia as análises
narrativas e a lista de pacientes (viram marcadores amarelos [PREENCHER]) e imprime os textos
com números que continuam fixos, para revisão. Rode de novo se o modelo-base mudar:

    python scripts/montar_template.py [base.docx] [saida.docx]
"""

import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from utils import sem_acento

DIAS_MAX = 31  # colunas de dia das tabelas diárias; meses mais curtos deixam as últimas em branco

RAIZ = Path(__file__).resolve().parent.parent
BASE = RAIZ / "templates" / "junho_convertido.docx"
SAIDA = RAIZ / "templates" / "modelo_relatorio.docx"

_MES_ANO = re.compile(r"(JANEIRO|FEVEREIRO|MAR[ÇC]O|ABRIL|MAIO|JUNHO|JULHO|AGOSTO|SETEMBRO|OUTUBRO|NOVEMBRO|DEZEMBRO)/(\d{4})")

VM = "taxa_de_utilizacao_de_ventilacao_mecanica"
# título compacto da tabela -> (variável, filtro). Ordem importa: mais específico primeiro.
TABELAS = [
    ("SAIDAS", "saidas_hospitalares__n_de_saidas", "num"),
    ("TAXADEOCUPACAOSEMOSPACIENTESCRONICOS", "taxa_de_ocupacao_sem_os_cronicos", "pct"),
    ("TAXADEOCUPACAO", "taxa_de_ocupacao_global", "pct"),
    ("MEDIADEPERMANENCIADIASSEMOSPACIENTESCRONICOS", "media_de_permanencia_sem_os_cronicos", "dec"),
    ("MEDIADEPERMANENCIADIASGLOBAL", "media_de_permanencia", "dec"),
    ("TAXADEMORTALIDADEINSTITUCIONAL", "taxa_de_mortalidade_em_uti", "pct"),
    ("100DEINTERCONSULTASNEURO", "atender_100_das_solicitacoes_de_interconsultas_em_neuro", "pct"),
    ("TAXADEREINTERNACAOEM24HORAS", "taxa_de_reinternacao_em_24h", "pct"),
    ("PNMASSOCIADAAVM", "pnm_associada_a_vm_pav", "pct"),
    ("TAXADEUTILIZACAODEVENTILACAOMECAN", VM, "pct"),
    ("INCIDENCIADEEXTUBACAOACIDENTAL", "incidencia_de_extubacao_acidental", "pct"),
    ("INCIDENCIADESAIDANAOPLANEJADADESONDA", "incidencia_de_saida_nao_planejada_de_sonda_oro_nasogastroenteral", "pct"),
    ("DENSIDADEDEINCIDENCIADEINFECCAOPRIMARIA", "densidade_de_incidencia_ipcs", "dec"),
    ("INCIDENCIADEFLEBITE", "incidencia_de_flebite", "pct"),
    ("INCIDENCIADEPERDADECATETERCENTRALDEINSERCAO", "incidencia_de_perda_de_cateter_central_insercao_periferica", "pct"),
    ("INCIDENCIADEPERDADECATETERCENTRAL", "incidencia_de_perda_de_cateter_central", "pct"),
    ("TAXADEUTILIZACAODECATETERVENOSOCENTRAL", "taxa_de_utilizacao_de_cateter_venoso_central_cvc", "pct"),
    ("INCIDENCIADEERROQUASEERRO", "incidencia_de_erro_quase_erro_de_medicacao", "pct"),
    ("DENSIDADEDEINCIDENCIADEINFECCOESDOTRATOURINARIO", "densidade_de_incidencia_itu", "dec"),
    ("INCIDENCIADEQUEDA", "incidencia_de_queda_de_paciente", "pct"),
    ("INDICEDEULCERA", "indice_de_ulcera_por_pressao", "pct"),
    ("ADESAOASMETAS", "adesao_as_metas_de_identificacao_do_paciente", "pct"),
    ("TAXADEADESAOAOPROTOCOLO", "taxa_de_adesao_ao_protocolo_de_hm_momentos", "pct"),
    ("EVENTOSENTINELA", "evento_sentinela_0", "dec"),
    ("100DASSOLICITACOESDETRANSPORTE", "atender_100_das_solicitacoes_de_transporte_inter_e_intrahospitalar", "pct"),
    ("FONOAUDIOLOGOTESTEDETRIAGEM", "fonoaudiologo_teste_triagem_auditiva_universal_teste_da_orelhinha_100_dos_recem_nascidos", "pct"),
    ("FONOAUDIOLOGOTESTEDELINGUINHA", "fonoaudiologo_teste_linguinha_em_100_dos_nascidos", "pct"),
]
# Coluna "Variação %": só a da reinternação tem fórmula clara (realizado / meta de 1%).
VARIACAO = {"TAXADEREINTERNACAOEM24HORAS": "{{ (taxa_de_reinternacao_em_24h / 0.01 * 100) | dec(0) }}%"}

_I = "indicador"


def _v(chave, filtro="num"):
    return "{{ " + chave + " | " + filtro + " }}"


def _t(base, chave):
    return _v(f"{base}__{chave}")


# Parágrafos de base de cálculo: (início do texto original, [textos novos em ordem de aparição]).
PARAGRAFOS = [
    ("Para base de cálculo da taxa de ocupação", [
        "Para base de cálculo da taxa de ocupação consideramos o quantitativo de "
        + _t("taxa_de_ocupacao_sem_os_cronicos", "n_de_pacientes_dia") + " pacientes/dia X "
        + _t("taxa_de_ocupacao_sem_os_cronicos", "n_de_leito_dias") + " leito/dia."]),
    ("Para base de cálculo da permanência", [
        "Para base de cálculo da permanência consideramos o quantitativo de "
        + _t("taxa_de_ocupacao_global", "n_de_pacientes_dia") + " pacientes/dia para todas as UTIs X "
        + _t("taxa_de_ocupacao_global", "n_de_leito_dias") + " leitos-dia"]),
    ("N.DE PACIENTE-DIA =", [
        "N.DE PACIENTE-DIA = " + _t("media_de_permanencia_sem_os_cronicos", "n_de_paciente_dia")
        + " / N. DE SAÍDAS = " + _t("media_de_permanencia_sem_os_cronicos", "n_de_saidas"),
        "N.DE PACIENTE-DIA = " + _t("media_de_permanencia", "n_de_paciente_dia")
        + " / N. DE SAÍDAS = " + _t("media_de_permanencia", "n_de_saidas")]),
    ("Base para cálculo", [
        "Base para cálculo " + _t("taxa_de_mortalidade_em_uti", "n_de_obitos")
        + " óbitos (mais de 24 horas) e " + _t("taxa_de_mortalidade_em_uti", "n_total_de_saidas") + " saídas"]),
    ("RESPOSTAS DENTRO DO PRAZO AS IC", [
        "RESPOSTAS DENTRO DO PRAZO AS IC = "
        + _t("atender_100_das_solicitacoes_de_interconsultas_em_neuro", "respostas_dentro_do_prazo_as_ic")
        + " / N. DE SOLICITAÇÕES REALIZADAS = "
        + _t("atender_100_das_solicitacoes_de_interconsultas_em_neuro", "n_de_solicitacoes_realizadas")]),
    ("N. DE REINTERNAÇÃO", [
        "N. DE REINTERNAÇÃO = " + _t("taxa_de_reinternacao_em_24h", "n_de_reinternacao")
        + " / N. TOTAL DE SAÍDAS = " + _t("taxa_de_reinternacao_em_24h", "n_total_de_saidas")]),
    ("N. DE PAV", [
        "N. DE PAV = " + _t("pnm_associada_a_vm_pav", "n_de_pav")
        + " / N. DE PACIENTES DIA EM VM = " + _t("pnm_associada_a_vm_pav", "n_de_pacientes_dia_em_vm")]),
    ("N. DE PACIENTES DIA EM VM", [
        "N. DE PACIENTES DIA EM VM = " + _t(VM, "n_de_pacientes_dia_em_vm")
        + " / N. TOTAL DE PACIENTE DIA = " + _t(VM, "n_total_de_paciente_dia")]),
    ("N. DE EXTUBAÇÃO NÃO PLANEJADA", [
        "N. DE EXTUBAÇÃO NÃO PLANEJADA = " + _t("incidencia_de_extubacao_acidental", "n_de_extubacao_nao_planejada")
        + " / N. DE PACIENTES ENTUBADOS/DIA = " + _t("incidencia_de_extubacao_acidental", "n_de_pacientes_entubados_dia")]),
    ("SAÍDA NÃO PLANEJADA DE SONDA", [
        "SAÍDA NÃO PLANEJADA DE SONDA ORO/NASOGASTROENTERAL = "
        + _t("incidencia_de_saida_nao_planejada_de_sonda_oro_nasogastroenteral", "saida_nao_planejada_de_sonda_oro_nasogastroenteral")
        + " / NÚMERO DE PACIENTES COM SONDA ORO/NASOGASTROENTERAL = "
        + _t("incidencia_de_saida_nao_planejada_de_sonda_oro_nasogastroenteral", "numero_de_pacientes_com_sonda_oro_nasogastroenteral")]),
    ("N. DE IPCS", [
        "N. DE IPCS = " + _t("densidade_de_incidencia_ipcs", "n_de_ipcs")
        + " / N. DE PACIENTE CVC – DIA = " + _t("densidade_de_incidencia_ipcs", "n_de_paciente_cvc_dia")]),
    ("N. DE FLEBITES", [
        "N. DE FLEBITES = " + _t("incidencia_de_flebite", "n_de_flebites")
        + " / N. DE PACIENTE - DIA COM CATETER VENOSO PERIFERICO = "
        + _t("incidencia_de_flebite", "n_de_paciente_dia_com_cateter_venoso_periferico")]),
    ("N. DE PERDA DE CATETER CENTRAL DE INSERÇÃO PERIFÉRICA", [
        "N. DE PERDA DE CATETER CENTRAL DE INSERÇÃO PERIFÉRICA = "
        + _t("incidencia_de_perda_de_cateter_central_insercao_periferica", "n_de_perda_de_cateter_central_de_insercao_periferica")
        + " / N. DE ACESSOS PERIFÉRICOS = "
        + _t("incidencia_de_perda_de_cateter_central_insercao_periferica", "n_de_paciente_dia_com_picc"),
        "N. DE PERDA DE CATETER CENTRAL DE INSERÇÃO PERIFÉRICA = "
        + _t("incidencia_de_perda_de_cateter_central", "n_de_perda_de_cateter_central")
        + " / N. DE PACIENTE - DIA COM CVC = "
        + _t("incidencia_de_perda_de_cateter_central", "n_de_paciente_dia_com_cvc")]),
    ("N. DE DIA DE CATETER CENTRAL", [
        "N. DE DIA DE CATETER CENTRAL = "
        + _t("taxa_de_utilizacao_de_cateter_venoso_central_cvc", "n_de_dia_de_cateter_central")
        + " / N. DE PACIENTE DIA = " + _t("taxa_de_utilizacao_de_cateter_venoso_central_cvc", "n_de_paciente_dia")]),
    ("N. DE ERROS RELACIONADOS", [
        "N. DE ERROS RELACIONADOS A MEDICAÇÃO/ FALHAS MEDICAÇÃO = "
        + _t("incidencia_de_erro_quase_erro_de_medicacao", "n_de_erros_relacionados_a_medicacao_falhas_medicacao")
        + " / N. DE PACIENTE DIA = " + _t("incidencia_de_erro_quase_erro_de_medicacao", "n_de_paciente_dia")]),
    ("N. DE CASOS NOVOS DE ITU", [
        "N. DE CASOS NOVOS DE ITU = " + _t("densidade_de_incidencia_itu", "n_de_casos_novos_de_itu")
        + " / PACIENTE DIA COM SVD = " + _t("densidade_de_incidencia_itu", "paciente_dia_com_svd")]),
    ("N. DE INCIDÊNCIA DE QUEDA", [
        "N. DE INCIDÊNCIA DE QUEDA DE PACIENTES = "
        + _t("incidencia_de_queda_de_paciente", "n_de_incidencia_de_queda_de_pacientes")
        + " / N. DE PACIENTE DIA = " + _t("incidencia_de_queda_de_paciente", "n_de_paciente_dia")]),
    ("Relação entre o número de casos novos", [
        "Relação entre o número de casos novos de pacientes com úlcera por pressão em um determinado período e o "
        "número de pessoas expostas ao risco de adquirir úlcera por pressão no período, multiplicado por 1 - "
        "N. DE CASOS NOVOS DE PACIENTE COM ÚLCERA POR PRESSÃO = "
        + _t("indice_de_ulcera_por_pressao", "n_de_casos_novos_de_paciente_com_ulcera_por_pressao")
        + " / N. DE PESSOAS EXPOSTAS AO RISCO DE ADQUIRIR = "
        + _t("indice_de_ulcera_por_pressao", "n_de_pessoas_expostas_ao_risco_de_adquirir") + "."]),
    ("Razão entre o número de eventos sentinelas", [
        "Razão entre o número de eventos sentinelas (" + _t("evento_sentinela_0", "n_de_eventos_sentinelas")
        + ") e o número de paciente dia (" + _t("evento_sentinela_0", "n_de_paciente_dia") + ")"]),
    ("Nº transportes realizados", [
        "Nº transportes realizados / Nº de Solicitações x 100 – Solicitados "
        + _t("atender_100_das_solicitacoes_de_transporte_inter_e_intrahospitalar", "n_de_solicitacoes")
        + "/realizados: "
        + _t("atender_100_das_solicitacoes_de_transporte_inter_e_intrahospitalar", "n_de_transportes_realizados")
        + " transportes"]),
    ("Nº de recém-nascidos que foram testados", [
        "Nº de recém-nascidos que foram testados/ nº nascidos vivos – "
        + _t("fonoaudiologo_teste_triagem_auditiva_universal_teste_da_orelhinha_100_dos_recem_nascidos", "n_de_recem_nascidos_que_foram_testados")
        + " RN testados /"
        + _t("fonoaudiologo_teste_triagem_auditiva_universal_teste_da_orelhinha_100_dos_recem_nascidos", "n_de_nascidos_vivos")
        + " nascidos vivos",
        "Nº de recém-nascidos que foram testados/ nº nascidos vivos – "
        + _t("fonoaudiologo_teste_linguinha_em_100_dos_nascidos", "n_de_recem_nascidos_que_foram_testados")
        + " RN testados /" + _t("fonoaudiologo_teste_linguinha_em_100_dos_nascidos", "n_de_nascidos_vivos")]),
    ("Tivemos um total de", ["Tivemos um total de " + _v("n_de_nascidos_vivos") + " nascidos vivos no HX."]),
    ("Taxa de Utilização de VM total", [
        "Taxa de Utilização de VM total: Considerando os números globais, tivemos uma taxa de "
        + _v(VM, "pct(1)") + ". [PREENCHER: taxa excluindo pacientes crônicos e análise]"]),
]

# Linhas das tabelas diárias de ventilação/crônicos: rótulo compacto -> variável da série.
LINHAS_DIARIAS = {
    "NDEPACIENTESDIAEMVM": f"{VM}__n_de_pacientes_dia_em_vm",
    "NTOTALDEPACIENTEDIA": f"{VM}__n_total_de_paciente_dia",
    "PACIENTEDIACRONICOS": "paciente_dia_cronicos",
}
SETORES_TESTES = {"UTIPED": "uti_ped", "UTINEO": "uti_neo", "PEDIATRIA": "pediatria",
                  "MATERNIDADE": "maternidade", "AMBULATORIO": "ambulatorio", "TOTAL": "total"}

PLACEHOLDER_JUSTIFICATIVA = "[PREENCHER: justificativa / análise do indicador]"


def _compacto(texto):
    return re.sub(r"[^A-Z0-9]", "", sem_acento(texto).upper())


def _set_par(p, texto, destaque=False):
    """Troca o texto do parágrafo mantendo a formatação do primeiro trecho."""
    runs = p.runs
    if not runs:
        runs = [p.add_run("")]
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if destaque or "[PREENCHER" in texto:
        runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def _set_cell(cell, texto):
    paragrafos = cell.paragraphs
    for extra in paragrafos[1:]:
        extra._p.getparent().remove(extra._p)
    _set_par(paragrafos[0], texto)


def _remover(el):
    el.getparent().remove(el)


def _largura(el):
    """Largura em dxa de um <w:gridCol> ou <w:tc> (None se ausente ou em outra unidade)."""
    if el.tag == qn("w:gridCol"):
        return int(el.get(qn("w:w"))) if el.get(qn("w:w")) else None
    tcw = el.tcPr.find(qn("w:tcW")) if el.tcPr is not None else None
    if tcw is None or tcw.get(qn("w:type"), "dxa") != "dxa" or not tcw.get(qn("w:w")):
        return None
    return int(tcw.get(qn("w:w")))


def _ajustar_largura(el, delta):
    atual = _largura(el)
    if atual is None:
        return
    alvo = el if el.tag == qn("w:gridCol") else el.tcPr.find(qn("w:tcW"))
    alvo.set(qn("w:w"), str(max(atual + delta, 0)))


def garantir_colunas_de_dia(tabela, dias=DIAS_MAX):
    """Tabela diária (rótulo | dia 1..N | TOTAL | taxa) com N < `dias` ganha as colunas que faltam.

    Cada coluna nova é cópia vazia da última coluna de dia, inserida logo antes de TOTAL. A largura
    sai da coluna do rótulo, para a tabela não crescer. Célula mesclada que cobre a última coluna de
    dia (título da tabela) ganha +1 no gridSpan. Idempotente: com `dias` colunas, nada muda.
    Retorna quantas colunas foram adicionadas.
    """
    tbl = tabela._tbl
    adicionadas = 0
    while len(tbl.tblGrid.gridCol_lst) - 3 < dias:
        cols = tbl.tblGrid.gridCol_lst
        pos = len(cols) - 2  # índice de grade da coluna nova (= posição atual de TOTAL)
        ultima_dia = cols[pos - 1]
        largura = _largura(ultima_dia) or 0
        ultima_dia.addnext(deepcopy(ultima_dia))
        _ajustar_largura(cols[0], -largura)
        for tr in tbl.tr_lst:
            inicio = 0
            for tc in list(tr.tc_lst):
                span = tc.grid_span
                fim = inicio + span  # exclusivo
                if span == 1 and inicio == pos - 1:  # última célula de dia: duplica, vazia
                    nova = deepcopy(tc)
                    for t in nova.iter(qn("w:t")):
                        t.text = ""
                    tc.addnext(nova)
                elif span > 1 and inicio < pos <= fim:  # mesclada sobre a última coluna de dia
                    tc.grid_span = span + 1
                elif inicio == 0 and span == 1:  # rótulo: cede a largura da coluna nova
                    _ajustar_largura(tc, -largura)
                inicio = fim
        adicionadas += 1
    return adicionadas


def montar(base=BASE, saida=SAIDA):
    doc = Document(base)
    elementos = list(doc.element.body.iterchildren())
    protegidos = set()

    # 1) Mês/ano no texto corrido (capa e introdução).
    # O mês/ano do documento-base é detectado (ex.: "JUNHO/2026"), não fixado no código.
    for p in doc.paragraphs:
        if _MES_ANO.search(p.text):
            _set_par(p, _MES_ANO.sub("{{ mes }}/{{ ano }}", p.text))

    # 2) Parágrafos de base de cálculo, por ordem de aparição.
    contadores = {inicio: 0 for inicio, _ in PARAGRAFOS}
    for el in elementos:
        if not el.tag.endswith("}p"):
            continue
        p = Paragraph(el, doc)
        texto = p.text.strip()
        for inicio, novos in PARAGRAFOS:
            if texto.startswith(inicio) and contadores[inicio] < len(novos):
                _set_par(p, novos[contadores[inicio]])
                contadores[inicio] += 1
                protegidos.add(el)
                break
    esperado = {i: len(n) for i, n in PARAGRAFOS}
    faltou = [i for i, c in contadores.items() if c != esperado[i]]
    if faltou:
        print("[aviso] parágrafos de base não encontrados no modelo-base:", faltou)

    # 3) Células "Realizado" das tabelas de indicador.
    sem_mapa = []
    for el in elementos:
        if not el.tag.endswith("}tbl"):
            continue
        tabela = Table(el, doc)
        titulo = _compacto(tabela.rows[0].cells[1].text) if len(tabela.rows[0].cells) > 1 else ""
        regra = next((r for r in TABELAS if titulo.startswith(r[0])), None)
        cab = next((r for r in tabela.rows[:-1] if any(c.text.strip() == "Realizado" for c in r.cells)), None)
        if regra and cab:
            idx = [c.text.strip() for c in cab.cells].index("Realizado")
            _set_cell(tabela.rows[-1].cells[idx], _v(regra[1], regra[2]))
            textos = [c.text.strip() for c in cab.cells]
            if "Variação %" in textos:
                var = VARIACAO.get(regra[0], "[PREENCHER: variação %]")
                _set_cell(tabela.rows[-1].cells[textos.index("Variação %")], var)
        elif cab:
            sem_mapa.append(tabela.rows[0].cells[1].text.strip())
    if sem_mapa:
        print("[aviso] tabelas com 'Realizado' sem mapeamento:", sem_mapa)

    # 4) Tabelas diárias (ventilação / crônicos) e tabela de testes.
    for el in elementos:
        if not el.tag.endswith("}tbl"):
            continue
        tabela = Table(el, doc)
        if len(tabela.columns) >= 33:
            if garantir_colunas_de_dia(tabela):  # o modelo-base (junho) tem 30 dias; o relatório precisa de 31
                print(f"[info] tabela diária ampliada para {DIAS_MAX} colunas de dia.")
            diaria_cronicos = False
            for r in tabela.rows[1:]:
                chave = LINHAS_DIARIAS.get(_compacto(r.cells[0].text))
                if not chave:
                    continue
                diaria_cronicos |= chave == "paciente_dia_cronicos"
                n = len(r.cells)
                for i in range(1, n - 2):
                    _set_cell(r.cells[i], "{{ serie__" + chave + " | dia(" + str(i) + ") }}")
                _set_cell(r.cells[n - 2], _v(chave))
                _set_cell(r.cells[n - 1], "[PREENCHER]" if diaria_cronicos else _v(VM, "pct"))
            n = len(tabela.rows[0].cells)
            _set_cell(tabela.rows[0].cells[n - 1], "[PREENCHER]" if diaria_cronicos else _v(VM, "pct"))
        elif tabela.rows[0].cells[0].text.strip() == "SETOR":
            for r in tabela.rows[1:]:
                setor = SETORES_TESTES.get(_compacto(r.cells[0].text))
                if setor:
                    _set_cell(r.cells[1], _v(f"teste_orelhinha_{setor}"))
                    _set_cell(r.cells[2], _v(f"teste_linguinha_{setor}"))
        elif tabela.rows[0].cells[0].text.strip() == "DATA ATENDIMENTO":
            for r in tabela.rows[2:]:
                _remover(r._tr)  # lista de pacientes: dado pessoal, preencher à mão
            for c in tabela.rows[1].cells:
                _set_cell(c, "[PREENCHER]")

    # 5) Análises narrativas: após "Justificativa / análise do indicador" até o próximo título/tabela.
    em_justificativa, primeiro = False, False
    antes_do_corpo = True  # texto solto antes do primeiro Heading 1 (sumário) também é removido
    for el in elementos:
        if el.tag.endswith("}tbl"):
            em_justificativa = False
            continue
        if not el.tag.endswith("}p"):
            continue
        p = Paragraph(el, doc)
        texto = p.text.strip()
        estilo = p.style.name if p.style is not None else ""
        if estilo.startswith("Heading 1") and texto:
            antes_do_corpo = False
        if antes_do_corpo and texto.startswith("- "):
            _remover(el)
            continue
        if texto == "Justificativa / análise do indicador":
            em_justificativa, primeiro = True, True
            continue
        # Há análises digitadas com estilo de título ("- Devido a..."): só é título se não começa com "-".
        if estilo.startswith("Heading") and texto and not (em_justificativa and texto.startswith("-")):
            em_justificativa = False
            continue
        if em_justificativa and texto and el not in protegidos:
            if primeiro:
                _set_par(p, PLACEHOLDER_JUSTIFICATIVA)
                primeiro = False
            else:
                _remover(el)
        elif em_justificativa and texto and el in protegidos:
            em_justificativa = False  # parágrafo de base de um novo indicador

    Path(saida).parent.mkdir(parents=True, exist_ok=True)
    doc.save(saida)
    return saida


def revisar(caminho):
    """Lista textos com dígitos que não têm variável: o que ainda está fixo no modelo."""
    doc = Document(caminho)
    fixos = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.search(r"\d", t) and "{{" not in t and "\t" not in t:
            fixos.append(t[:110])
    for tabela in doc.tables:
        for r in tabela.rows:
            for c in r.cells:
                t = c.text.strip()
                if re.search(r"\d", t) and "{{" not in t and t not in fixos and len(t) > 2:
                    fixos.append("[tabela] " + t[:60])
    return list(dict.fromkeys(fixos))


if __name__ == "__main__":
    base = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE
    saida = Path(sys.argv[2]) if len(sys.argv) > 2 else SAIDA
    montar(base, saida)
    print("Modelo salvo em:", saida)
    print("--- Textos com número ainda FIXOS (conferir: contrato/meta = ok; resto = ajustar) ---")
    for t in revisar(saida):
        print("  ", t)
